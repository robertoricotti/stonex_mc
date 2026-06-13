import os
import math
import uuid
import csv
import shutil
import sys
from datetime import datetime, timezone

import tkinter as tk
from tkinter import filedialog, messagebox

from tkinterdnd2 import DND_FILES, TkinterDnD
import openpyxl
import xml.etree.ElementTree as ET


# ----------------------------
# Namespaces
# ----------------------------
NS_DRP = "http://www.iredes.org/xml/DrillRig"
NS_IR  = "http://www.iredes.org/xml"
NS_JET = "urn:custom:jet:treatment:v1"   # DEVE combaciare con Android parser

ET.register_namespace("", NS_DRP)
ET.register_namespace("IR", NS_IR)
ET.register_namespace("JET", NS_JET)


# ----------------------------
# Utility
# ----------------------------
def _is_empty(v) -> bool:
    if v is None:
        return True
    if isinstance(v, str) and v.strip() == "":
        return True
    return False


def _clean_cell(v):
    if v is None:
        return None
    s = str(v).strip()
    if len(s) >= 2 and s[0] == '"' and s[-1] == '"':
        s = s[1:-1].strip()
    s = s.lstrip("\ufeff").strip()
    return None if s == "" else s


def desktop_path() -> str:
    home = os.path.expanduser("~")
    cand = os.path.join(home, "Desktop")
    return cand if os.path.isdir(cand) else home


def app_dir() -> str:
    """
    Ritorna la cartella base dell'app.
    - In sviluppo: la cartella dove sta lo script .py
    - In exe (PyInstaller): la cartella dell'eseguibile
    """
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


def bundled_dir() -> str:
    """
    Cartella temporanea usata da PyInstaller --onefile per i file inclusi
    con --add-data. In sviluppo coincide con app_dir().
    """
    return getattr(sys, "_MEIPASS", app_dir())


def resource_path(relative: str) -> str:
    """
    Restituisce un path valido per risorse esterne come templates e icone.

    Cerca in questo ordine:
    1) cartella temporanea PyInstaller (_MEIPASS), se l'exe è stato creato con --add-data
    2) cartella dell'eseguibile / dello script
    3) cartella corrente da cui viene lanciata l'app
    4) cartella dello script sorgente, se disponibile

    Questo rende l'app più robusta sia in sviluppo sia come .exe.
    """
    candidates = [
        os.path.join(bundled_dir(), relative),
        os.path.join(app_dir(), relative),
        os.path.join(os.getcwd(), relative),
    ]

    try:
        script_dir = os.path.dirname(os.path.abspath(__file__))
        candidates.append(os.path.join(script_dir, relative))
    except Exception:
        pass

    for candidate in candidates:
        if os.path.exists(candidate):
            return candidate

    # Fallback: torna al path PyInstaller, così l'errore resta leggibile.
    return os.path.join(bundled_dir(), relative)


TEMPLATE_FILES = {
    "JET": ("templates/Jet_Template.xlsx", "Jet_Template.xlsx"),
    "ROCK": ("templates/Rock_Template.xlsx", "Rock_Template.xlsx"),
    "SOLAR": ("templates/Solar_Template.xlsx", "Solar_Template.xlsx"),
}


def copy_template_if_missing(src_relative_path: str, dst_folder: str, dst_name: str) -> bool:
    """
    Copia un template nella cartella progetto solo se manca.
    Ritorna True se il template è presente/copiato, False se la sorgente non è stata trovata.
    """
    os.makedirs(dst_folder, exist_ok=True)
    dst = os.path.join(dst_folder, dst_name)
    if os.path.exists(dst):
        return True  # non sovrascrivere

    src_path = resource_path(src_relative_path)
    if os.path.exists(src_path):
        shutil.copyfile(src_path, dst)
        return True

    return False


def ensure_project_folders() -> dict:
    base = os.path.join(desktop_path(), "STX Hole Converter")
    rock = os.path.join(base, "ROCK Proj")
    jet  = os.path.join(base, "JET Proj")
    solar= os.path.join(base, "SOLAR Proj")
    generic = os.path.join(base, "GENERIC CSV Proj")

    os.makedirs(rock, exist_ok=True)
    os.makedirs(jet, exist_ok=True)
    os.makedirs(solar, exist_ok=True)
    os.makedirs(generic, exist_ok=True)

    missing_templates = []

    targets = {
        "JET": jet,
        "ROCK": rock,
        "SOLAR": solar,
    }

    for mode, dst_folder in targets.items():
        src_rel, dst_name = TEMPLATE_FILES[mode]
        if not copy_template_if_missing(src_rel, dst_folder, dst_name):
            missing_templates.append(dst_name)

    return {
        "base": base,
        "ROCK": rock,
        "JET": jet,
        "SOLAR": solar,
        "AUTO": generic,
        "__missing_templates": missing_templates,
    }



def open_in_explorer(path: str):
    try:
        os.startfile(path)
    except Exception:
        pass


def compute_geometry(xh, yh, zh, xe, ye, ze):
    dx = xe - xh
    dy = ye - yh
    dz = ze - zh
    horiz = math.hypot(dx, dy)
    length = math.sqrt(dx * dx + dy * dy + dz * dz)
    bearing = math.degrees(math.atan2(dx, dy))
    if bearing < 0:
        bearing += 360.0
    vert = abs(dz)
    tilt = math.degrees(math.atan2(horiz, vert)) if not (horiz == 0 and vert == 0) else 0.0
    return {
        "dx": dx, "dy": dy, "dz": dz,
        "horiz": horiz, "length": length,
        "bearing_deg": bearing,
        "tilt_deg": tilt,
    }


def parse_float(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().strip('"').strip()
    if s == "":
        return None
    # Supporta sia decimale con virgola sia numeri con separatore migliaia.
    if "," in s and "." in s:
        # 1.234,56 -> 1234.56 ; 1,234.56 -> 1234.56
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    else:
        s = s.replace(",", ".")
    try:
        return float(s)
    except Exception:
        return None


def format_point_id(v, fallback_index: int) -> str:
    if _is_empty(v):
        return str(fallback_index)
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


# ----------------------------
# Readers (Excel/CSV)
# ----------------------------
def read_table_from_excel(path: str):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    headers = [("" if h is None else str(h).strip().lstrip("\ufeff")) for h in headers]

    rows = []
    for r in range(2, ws.max_row + 1):
        row = {}
        empty_row = True
        for c, h in enumerate(headers, start=1):
            if h == "":
                continue
            v = ws.cell(r, c).value
            if not _is_empty(v):
                empty_row = False
            row[h] = v
        if not empty_row:
            rows.append(row)
    return headers, rows


def _read_text_lines(path: str):
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            with open(path, "r", encoding=enc, newline="") as f:
                return f.read().splitlines()
        except UnicodeDecodeError:
            continue
    with open(path, "r", encoding="utf-8", errors="replace", newline="") as f:
        return f.read().splitlines()


def _strip_outer_quotes(line: str) -> str:
    s = line.strip()
    if len(s) >= 2 and s[0] == '"' and s[-1] == '"':
        return s[1:-1]
    return s


def _choose_delimiter(lines):
    candidates = [";", ",", "\t"]
    sample = [_strip_outer_quotes(l) for l in lines if l.strip()][:50]
    if not sample:
        return ","
    scores = {d: sum(max(0, len(row.split(d)) - 1) for row in sample) for d in candidates}
    return max(scores, key=scores.get)


def read_table_from_csv(path: str):
    """
    Reader CSV robusto:
    - header opzionale
    - delimitatori ; , TAB
    - supporta file in cui ogni riga è quotata interamente: "NUMBER;X;Y;Z"
    - se non c'è header, mette le colonne in __cols per il parser generico
    """
    lines = _read_text_lines(path)
    delim = _choose_delimiter(lines)
    rows_as_lists = []
    for line in lines:
        if not line.strip():
            continue
        line = _strip_outer_quotes(line)
        parsed = next(csv.reader([line], delimiter=delim, quotechar='"'))
        parsed = [_clean_cell(x) for x in parsed]
        if any(not _is_empty(x) for x in parsed):
            rows_as_lists.append(parsed)

    if not rows_as_lists:
        raise ValueError("CSV vuoto o non leggibile.")

    first = rows_as_lists[0]
    has_header = looks_like_header(first)

    if has_header:
        headers = [("" if h is None else str(h).strip().lstrip("\ufeff")) for h in first]
        data = rows_as_lists[1:]
        rows = []
        for rr in data:
            row = {}
            empty_row = True
            for i, h in enumerate(headers):
                if h == "":
                    continue
                v = rr[i] if i < len(rr) else None
                if not _is_empty(v):
                    empty_row = False
                row[h] = v
            if not empty_row:
                rows.append(row)
        return headers, rows

    rows = []
    for rr in rows_as_lists:
        if any(not _is_empty(x) for x in rr):
            rows.append({"__cols": rr})
    max_cols = max(len(r["__cols"]) for r in rows) if rows else 0
    return [f"COL{i}" for i in range(1, max_cols + 1)], rows


# ----------------------------
# Header mapping per CSV generici
# ----------------------------
def _norm_header(h: str) -> str:
    return "".join(ch for ch in str(h).strip().lstrip("\ufeff").lower() if ch.isalnum())


HEADER_ALIASES = {
    "P": {"p", "pt", "pnt", "ptnr", "point", "pointid", "pointname", "number", "num", "id", "name", "hole", "holeid", "holenumber"},
    "E": {"e", "east", "easting", "x", "xcoord", "coordx", "pointx"},
    "N": {"n", "north", "northing", "y", "ycoord", "coordy", "pointy"},
    "Z": {"z", "el", "elev", "elevation", "quota", "quote", "rl", "pointz", "coordz", "zcoord"},
    "D": {"d", "desc", "descr", "description", "comment", "comments", "note", "notes", "code", "color", "colour", "layer"},
}


def header_role(h: str):
    n = _norm_header(h)
    for role, aliases in HEADER_ALIASES.items():
        if n in aliases:
            return role
    return None


def looks_like_header(cells) -> bool:
    roles = [header_role(c) for c in cells]
    if any(r in ("E", "N", "Z") for r in roles):
        return True
    # Header test conservativo: evita di scambiare una riga dati con ID testuale per header.
    numeric_count = sum(1 for c in cells if parse_float(c) is not None)
    text_count = sum(1 for c in cells if parse_float(c) is None and not _is_empty(c))
    return text_count >= 2 and numeric_count == 0


def mapped_row_values(row: dict):
    values = {"P": None, "E": None, "N": None, "Z": None, "D": None}
    for k, v in row.items():
        if k == "__cols":
            continue
        role = header_role(k)
        if role and values[role] is None:
            values[role] = v
    return values


def _coord_order_from_two_values(a, b, requested_order="AUTO"):
    """
    Restituisce 'EN' o 'NE' per due colonne coordinate.
    In AUTO usa una euristica semplice: se una coordinata è molto più grande dell'altra,
    spesso quella più grande è Northing. Se non è chiaro, PENZD/EN è il default.
    """
    req = (requested_order or "AUTO").upper().replace(" ", "")
    if req in ("PNEZD", "NEZD", "NE", "YXZ"):
        return "NE"
    if req in ("PENZD", "ENZD", "EN", "XYZ"):
        return "EN"

    fa = abs(parse_float(a) or 0.0)
    fb = abs(parse_float(b) or 0.0)
    if fa > 2_000_000 and fb < 2_000_000:
        return "NE"
    if fb > 2_000_000 and fa < 2_000_000:
        return "EN"
    return "EN"


def _generic_values_from_cols(cols, idx: int, requested_order="AUTO"):
    cols = list(cols)
    # P presente se la prima colonna non è numerica, oppure se ci sono 4/5 colonne e la prima sembra ID.
    first_num = parse_float(cols[0]) if len(cols) > 0 else None
    p_present = False
    if len(cols) >= 5:
        p_present = True
    elif len(cols) >= 3 and first_num is None:
        p_present = True

    offset = 1 if p_present else 0
    p = cols[0] if p_present else None

    if len(cols) - offset < 2:
        return None

    c1 = cols[offset]
    c2 = cols[offset + 1]
    order = _coord_order_from_two_values(c1, c2, requested_order)
    if order == "NE":
        n, e = c1, c2
    else:
        e, n = c1, c2

    z = None
    d = None
    rest = cols[offset + 2:]
    if rest:
        if parse_float(rest[0]) is not None:
            z = rest[0]
            if len(rest) > 1:
                d = rest[1]
        else:
            d = rest[0]
            if len(rest) > 1 and parse_float(rest[1]) is not None:
                z = rest[1]

    return {"P": p, "E": e, "N": n, "Z": z, "D": d}


def is_known_template_headers(headers, mode=None):
    h = set(headers or [])
    if mode == "JET" or {"PtNr", "E Head", "N Head", "Z Head", "E End", "N End", "Z End"}.issubset(h):
        return "JET"
    if mode == "ROCK" or {"RowNr", "PtNr", "E Head", "N Head", "Z Head", "E End", "N End", "Z End"}.issubset(h):
        return "ROCK"
    if mode == "SOLAR" or {"PtNr", "E Head", "N Head", "Z Head"}.issubset(h):
        return "SOLAR"
    return None


# ----------------------------
# Canonical conversion per modalità
# ----------------------------
def rows_to_canonical(mode: str, rows: list, single_point_delta_m: float, headers=None, csv_order="AUTO"):
    """
    Ritorna lista di dict canonici:
    {
      id, rowId, description,
      headX,headY,headZ, endX,endY,endZ,
      diameter,
      treatment: { ... }   # solo per JET
    }

    Regole aggiunte:
    - CSV generic PNEZD/PENZD/ENZD/NEZD con header opzionale.
    - P mancante => numerazione progressiva.
    - Z mancante => 0.
    - D mancante/nulla => testo vuoto.
    - se c'è un solo punto: EndPoint verticale, Z_end = Z_start - distanza.
    """
    out = []
    mode = (mode or "AUTO").upper()
    template_mode = is_known_template_headers(headers or [], None if mode == "AUTO" else mode)

    if template_mode == "JET":
        for r in rows:
            pid = r.get("PtNr")
            if _is_empty(pid):
                continue
            xh = parse_float(r.get("E Head"))
            yh = parse_float(r.get("N Head"))
            zh = parse_float(r.get("Z Head"))
            xe = parse_float(r.get("E End"))
            ye = parse_float(r.get("N End"))
            ze = parse_float(r.get("Z End"))
            if None in (xh, yh, zh, xe, ye, ze):
                continue
            treatment_keys = [
                "pr_1","drlStart_1","drlStop_1", "pr_2","drlStart_2","drlStop_2",
                "pr_3","drlStart_3","drlStop_3", "pr_4","drlStart_4","drlStop_4",
                "pr_j_1","jetStart_1","jetStop_1", "pr_j_2","jetStart_2","jetStop_2",
                "pr_j_3","jetStart_3","jetStop_3", "pr_j_4","jetStart_4","jetStop_4",
            ]
            treatment = {k: str(r.get(k)).strip() for k in treatment_keys if not _is_empty(r.get(k))}
            out.append({
                "id": str(pid).strip(), "rowId": None, "description": None,
                "headX": xh, "headY": yh, "headZ": zh,
                "endX": xe, "endY": ye, "endZ": ze,
                "diameter": None, "treatment": treatment
            })
        return out

    if template_mode == "ROCK":
        for r in rows:
            rownr = r.get("RowNr")
            pid = r.get("PtNr")
            if _is_empty(pid) and _is_empty(rownr):
                continue
            xh = parse_float(r.get("E Head")); yh = parse_float(r.get("N Head")); zh = parse_float(r.get("Z Head"))
            xe = parse_float(r.get("E End"));  ye = parse_float(r.get("N End"));  ze = parse_float(r.get("Z End"))
            if None in (xh, yh, zh, xe, ye, ze):
                continue
            dia = parse_float(r.get("Diameter"))
            desc = r.get("Description")
            desc = "" if _is_empty(desc) else str(desc).strip()
            rid = None if _is_empty(rownr) else str(rownr).strip()
            pid_s = None if _is_empty(pid) else str(pid).strip()
            hole_name = f"{rid}.{pid_s}" if rid and pid_s else (pid_s or rid or str(len(out) + 1))
            out.append({
                "id": hole_name, "rowId": rid, "description": desc,
                "headX": xh, "headY": yh, "headZ": zh,
                "endX": xe, "endY": ye, "endZ": ze,
                "diameter": dia, "treatment": {}
            })
        return out

    if template_mode == "SOLAR":
        for r in rows:
            pid = r.get("PtNr")
            if _is_empty(pid):
                pid = len(out) + 1
            xh = parse_float(r.get("E Head")); yh = parse_float(r.get("N Head")); zh = parse_float(r.get("Z Head"))
            if None in (xh, yh):
                continue
            if zh is None:
                zh = 0.0
            desc = r.get("Description")
            desc = "" if _is_empty(desc) else str(desc).strip()
            out.append({
                "id": str(pid).strip(), "rowId": None, "description": desc,
                "headX": xh, "headY": yh, "headZ": zh,
                "endX": xh, "endY": yh, "endZ": zh - float(single_point_delta_m),
                "diameter": None, "treatment": {}
            })
        return out

    # Generic CSV: header oppure colonne libere.
    for i, r in enumerate(rows, start=1):
        vals = _generic_values_from_cols(r.get("__cols"), i, csv_order) if "__cols" in r else mapped_row_values(r)
        if vals is None:
            continue
        pid = format_point_id(vals.get("P"), i)
        xh = parse_float(vals.get("E"))
        yh = parse_float(vals.get("N"))
        zh = parse_float(vals.get("Z"))
        if None in (xh, yh):
            continue
        if zh is None:
            zh = 0.0
        desc = vals.get("D")
        desc = "" if _is_empty(desc) else str(desc).strip()
        out.append({
            "id": pid, "rowId": None, "description": desc,
            "headX": xh, "headY": yh, "headZ": zh,
            "endX": xh, "endY": yh, "endZ": zh - float(single_point_delta_m),
            "diameter": None, "treatment": {}
        })

    return out


# ----------------------------
# IREDES Writer
# ----------------------------
def build_iredes_doc(plan_name: str, holes_canon: list, drillbit_dia_mm: int):
    root = ET.Element(f"{{{NS_DRP}}}DRPPlan", {
        "DRPPlanVersion": "V 1.0",
        "IRVersion": "V 1.0",
        "IRDownwCompat": "V 1.0",
        "DRPPlanDownwCompat": "V 1.0",
    })

    genhead = ET.SubElement(root, f"{{{NS_IR}}}GenHead")
    ET.SubElement(genhead, f"{{{NS_IR}}}FileCreateDate").text = (
        datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    )
    irver = ET.SubElement(genhead, f"{{{NS_IR}}}IRversion", {"DownwCompat": "V 1.0"})
    irver.text = "V 1.0"

    ET.SubElement(root, f"{{{NS_IR}}}PlanId").text = uuid.uuid4().hex[:12]
    ET.SubElement(root, f"{{{NS_IR}}}PlanName").text = plan_name
    ET.SubElement(root, f"{{{NS_IR}}}Comment").text = ""
    ET.SubElement(root, f"{{{NS_IR}}}Project").text = ""
    ET.SubElement(root, f"{{{NS_IR}}}WorkOrder").text = ""

    drillplan = ET.SubElement(root, "DrillPlan")
    ET.SubElement(drillplan, "NumberOfHoles").text = str(len(holes_canon))

    hole_id = 1
    for hc in holes_canon:
        hole = ET.SubElement(drillplan, "Hole")
        ET.SubElement(hole, "HoleId").text = str(hole_id)
        ET.SubElement(hole, "HoleName").text = str(hc["id"])

        # Comment sempre presente: D mancante/nulla => stringa vuota.
        ET.SubElement(hole, "Comment").text = str(hc.get("description") or "")

        sp = ET.SubElement(hole, "StartPoint")
        ET.SubElement(sp, f"{{{NS_IR}}}PointX").text = f"{hc['headX']:.3f}"
        ET.SubElement(sp, f"{{{NS_IR}}}PointY").text = f"{hc['headY']:.3f}"
        ET.SubElement(sp, f"{{{NS_IR}}}PointZ").text = f"{hc['headZ']:.3f}"

        ep = ET.SubElement(hole, "EndPoint")
        ET.SubElement(ep, f"{{{NS_IR}}}PointX").text = f"{hc['endX']:.3f}"
        ET.SubElement(ep, f"{{{NS_IR}}}PointY").text = f"{hc['endY']:.3f}"
        ET.SubElement(ep, f"{{{NS_IR}}}PointZ").text = f"{hc['endZ']:.3f}"

        ET.SubElement(hole, "TypeOfHole").text = "Undefined"
        if hc.get("diameter") is not None:
            ET.SubElement(hole, "DrillBitDia").text = str(int(round(hc["diameter"])))
        else:
            ET.SubElement(hole, "DrillBitDia").text = str(int(drillbit_dia_mm))

        g = compute_geometry(hc["headX"], hc["headY"], hc["headZ"], hc["endX"], hc["endY"], hc["endZ"])
        jet_geom = ET.SubElement(hole, f"{{{NS_JET}}}Geometry")
        ET.SubElement(jet_geom, f"{{{NS_JET}}}Length").text = f"{g['length']:.3f}"
        ET.SubElement(jet_geom, f"{{{NS_JET}}}BearingDeg").text = f"{g['bearing_deg']:.3f}"
        ET.SubElement(jet_geom, f"{{{NS_JET}}}TiltDeg").text = f"{g['tilt_deg']:.3f}"
        ET.SubElement(jet_geom, f"{{{NS_JET}}}DeltaZ").text = f"{g['dz']:.3f}"

        treatment = hc.get("treatment") or {}
        if len(treatment) > 0:
            tr = ET.SubElement(hole, f"{{{NS_JET}}}Treatment")
            for k, v in treatment.items():
                tag = "".join(ch if (ch.isalnum() or ch in ["_", "-"]) else "_" for ch in k.strip())
                ET.SubElement(tr, f"{{{NS_JET}}}{tag}").text = str(v)
        hole_id += 1

    tree = ET.ElementTree(root)
    ET.indent(tree, space="    ", level=0)
    return tree


def unit_factor_to_meters(unit_name: str) -> float:
    u = (unit_name or "m").strip().lower()
    if u in ("m", "meter", "meters", "metre", "metres"):
        return 1.0
    if u in ("usfeet", "us foot", "us feet", "ft", "foot", "feet"):
        return 0.3048
    if u in ("ussurveyfeet", "us survey feet", "surveyfeet", "survey feet", "sft"):
        return 1200.0 / 3937.0
    raise ValueError(f"Unità non supportata: {unit_name}")


def convert_file(mode: str, in_path: str, single_point_distance: float, single_point_unit: str, drillbit_dia_mm: int, csv_order="AUTO"):
    ext = os.path.splitext(in_path)[1].lower()
    if ext in [".xlsx", ".xlsm"]:
        headers, rows = read_table_from_excel(in_path)
    elif ext == ".xls":
        raise ValueError("Formato .xls legacy non supportato da openpyxl. Salva il file come .xlsx o .csv.")
    elif ext == ".csv":
        headers, rows = read_table_from_csv(in_path)
    else:
        raise ValueError("Unsupported file Format. Use .xlsx/.xlsm o .csv")

    delta_m = float(single_point_distance) * unit_factor_to_meters(single_point_unit)
    holes_canon = rows_to_canonical(mode, rows, single_point_delta_m=delta_m, headers=headers, csv_order=csv_order)

    if len(holes_canon) == 0:
        raise ValueError("No valid Hole found (check column/values).")

    plan_name = os.path.splitext(os.path.basename(in_path))[0]
    tree = build_iredes_doc(plan_name, holes_canon, drillbit_dia_mm=drillbit_dia_mm)

    out_path = os.path.splitext(in_path)[0] + ".ird"
    tree.write(out_path, encoding="utf-8", xml_declaration=True)
    return out_path, len(holes_canon)


# ----------------------------
# GUI
# ----------------------------
class App(TkinterDnD.Tk):
    def __init__(self):
        super().__init__()
        self.title("STX Hole Converter")
        self.geometry("980x610")
        icon_path = resource_path("stx_icon.ico")
        if os.path.exists(icon_path):
            try:
                self.iconbitmap(icon_path)
            except Exception:
                pass

        self.folders = ensure_project_folders()
        self.mode = "AUTO"  # "AUTO" | "ROCK" | "JET" | "SOLAR"
        self.after(300, self.warn_if_templates_missing)

        top = tk.Frame(self)
        top.pack(fill="x", padx=12, pady=10)

        self.btn_auto = tk.Button(top, text="AUTO / GENERIC CSV", height=2, command=lambda: self.enter_mode("AUTO"))
        self.btn_rock = tk.Button(top, text="ROCK DRILL TEMPLATE", height=2, command=lambda: self.enter_mode("ROCK"))
        self.btn_jet  = tk.Button(top, text="JET GROUTING TEMPLATE", height=2, command=lambda: self.enter_mode("JET"))
        self.btn_solar= tk.Button(top, text="SOLAR FARM TEMPLATE", height=2, command=lambda: self.enter_mode("SOLAR"))

        self.btn_auto.pack(side="left", expand=True, fill="x", padx=4)
        self.btn_rock.pack(side="left", expand=True, fill="x", padx=4)
        self.btn_jet.pack(side="left",  expand=True, fill="x", padx=4)
        self.btn_solar.pack(side="left",expand=True, fill="x", padx=4)

        opts = tk.LabelFrame(self, text="Opzioni")
        opts.pack(fill="x", padx=12, pady=(0, 10))

        self.var_drillbit = tk.StringVar(value="102")
        self.var_single_distance = tk.StringVar(value="1.0")
        self.var_unit = tk.StringVar(value="m")
        self.var_csv_order = tk.StringVar(value="AUTO")

        tk.Label(opts, text="DrillBitDia default (mm):").grid(row=0, column=0, sticky="w", padx=8, pady=6)
        tk.Entry(opts, textvariable=self.var_drillbit, width=10).grid(row=0, column=1, sticky="w", padx=8, pady=6)

        tk.Label(opts, text="Start-End distance with only 1 point:").grid(row=0, column=2, sticky="w", padx=8, pady=6)
        tk.Entry(opts, textvariable=self.var_single_distance, width=10).grid(row=0, column=3, sticky="w", padx=8, pady=6)

        tk.Label(opts, text="Units:").grid(row=0, column=4, sticky="w", padx=8, pady=6)
        tk.OptionMenu(opts, self.var_unit, "m", "usFeet", "UsSurveyFeet").grid(row=0, column=5, sticky="w", padx=8, pady=6)

        tk.Label(opts, text="CSV without header:").grid(row=1, column=0, sticky="w", padx=8, pady=6)
        tk.OptionMenu(opts, self.var_csv_order, "AUTO", "PENZD", "PNEZD", "ENZD", "NEZD").grid(row=1, column=1, sticky="w", padx=8, pady=6)
        tk.Label(opts, text="AUTO tries to distinguish E/N vs N/E; if it's not clear, use PENZD.").grid(row=1, column=2, columnspan=4, sticky="w", padx=8, pady=6)

        self.drop = tk.Label(
            self,
            text="Mode: AUTO / GENERIC CSV\nDrag a file (xlsx/xlsm/csv) here\nor click to choose the file.",
            relief="groove",
            borderwidth=2,
            height=12,
            font=("Segoe UI", 11)
        )
        self.drop.pack(fill="both", expand=True, padx=12, pady=(0, 10))

        self.drop.drop_target_register(DND_FILES)
        self.drop.dnd_bind("<<Drop>>", self.on_drop)
        self.drop.bind("<Button-1>", self.on_click_pick)

        self.status = tk.Label(self, text="Pronto.", anchor="w")
        self.status.pack(fill="x", padx=12, pady=(0, 12))
        self.refresh_ui_state()

    def warn_if_templates_missing(self):
        missing = self.folders.get("__missing_templates", [])
        if not missing:
            return
        messagebox.showwarning(
            "Template mancanti",
            "Non ho trovato questi template:\n\n"
            + "\n".join(f"- {name}" for name in missing)
            + "\n\nPer includerli nell'exe, rigenera con:\n"
              'python -m PyInstaller --onefile --windowed --name stonex_converter '
              '--icon stx_icon.ico --add-data "templates;templates" --add-data "stx_icon.ico;." stx_hole_converter.py'
        )

    def refresh_ui_state(self):
        self.drop.config(
            text=(
                f"Mode: {self.mode}\n"
                "Drag the file here (xlsx/xlsm/csv) to convert it in .ird\n"
                "Known templates are detected automatically; generic CSV accepts PNEZD/PENZD with optional header.\n"
                "Output: same name, same folder."
            )
        )

    def set_status(self, msg: str):
        self.status.config(text=msg)
        self.update_idletasks()

    def enter_mode(self, mode: str):
        self.mode = mode
        self.folders = ensure_project_folders()
        open_in_explorer(self.folders.get(mode, self.folders["AUTO"]))
        self.set_status(f"Selected Mode: {mode}.")
        self.refresh_ui_state()

    def on_click_pick(self, event=None):
        path = filedialog.askopenfilename(
            initialdir=self.folders.get(self.mode, self.folders["AUTO"]),
            filetypes=[("Input files", "*.xlsx *.xlsm *.csv")]
        )
        if path:
            self.convert(path)

    def on_drop(self, event):
        files = self.tk.splitlist(event.data)
        if not files:
            return
        self.convert(files[0])

    def convert(self, in_path: str):
        try:
            if not os.path.isfile(in_path):
                raise FileNotFoundError("File not found.")
            drillbit = int(float(self.var_drillbit.get().strip().replace(",", ".")))
            single_distance = float(self.var_single_distance.get().strip().replace(",", "."))
            unit = self.var_unit.get()
            csv_order = self.var_csv_order.get()

            self.set_status(f"Converting ({self.mode}): {os.path.basename(in_path)} ...")
            out_path, n = convert_file(
                mode=self.mode,
                in_path=in_path,
                single_point_distance=single_distance,
                single_point_unit=unit,
                drillbit_dia_mm=drillbit,
                csv_order=csv_order,
            )
            self.set_status(f"OK: created {os.path.basename(out_path)} ({n} holes)")
            messagebox.showinfo("Completed", f"File Created:\n{out_path}\nHoles: {n}")
        except Exception as e:
            self.set_status("Conversion Error.")
            messagebox.showerror("Error", str(e))


if __name__ == "__main__":
    App().mainloop()
